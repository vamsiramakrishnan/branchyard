"""Tests for tools/console_fixtures.py: reopen every artifact and check the manifest."""

from __future__ import annotations

import csv
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "sdk" / "python"))

import branchyard_artifacts  # noqa: E402
import console_fixtures  # noqa: E402


def _available(name: str) -> bool:
    import importlib.util

    return importlib.util.find_spec(name) is not None


@unittest.skipUnless(
    all(_available(name) for name in ["docx", "openpyxl", "pptx", "jsonschema", "blake3"]),
    "Complete knowledge fixtures need the Python artifacts extra",
)
class ConsoleFixturesTest(unittest.TestCase):
    """Generate once per test method into a fresh scratch directory."""

    def setUp(self) -> None:
        self._temporary = tempfile.TemporaryDirectory(prefix="console-fixtures-test-")
        self.addCleanup(self._temporary.cleanup)
        self.scratch = Path(self._temporary.name)
        self.out = self.scratch / "bundle"
        self.manifest = console_fixtures.generate(self.out)

    def test_manifest_matches_files_on_disk(self) -> None:
        names = {entry["filename"] for entry in self.manifest["files"]}
        expected_names = {"ticket_backlog.csv", "backlog_metrics.xlsx", "backlog_report.docx", "backlog_briefing.pptx"}
        self.assertEqual(names, expected_names)
        for entry in self.manifest["files"]:
            data = (self.out / entry["filename"]).read_bytes()
            self.assertEqual(entry["bytes"], len(data))
            self.assertEqual(entry["sha256"], hashlib.sha256(data).hexdigest())
            self.assertIsInstance(entry["media_type"], str)
            self.assertIsInstance(entry["expected"], dict)
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), sorted([*names, "manifest.json"]))

    def test_csv_keeps_identifiers_and_inert_text(self) -> None:
        with (self.out / "ticket_backlog.csv").open(newline="", encoding="utf-8") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["ticket_id"], "00451")
        self.assertEqual(rows[1]["summary"], 'Export dialog drops quotes, "fields"')
        self.assertIn("=SUM(A1:A2)", rows[0]["notes"])
        self.assertIn("<script>", rows[1]["notes"])

    @unittest.skipUnless(_available("openpyxl"), "openpyxl is not installed")
    def test_xlsx_two_sheets_formula_and_chart(self) -> None:
        import openpyxl

        workbook = openpyxl.load_workbook(self.out / "backlog_metrics.xlsx")
        self.assertEqual(workbook.sheetnames, ["Backlog", "Summary"])
        self.assertEqual(workbook["Backlog"]["A2"].value, "00451")
        summary = workbook["Summary"]
        self.assertEqual(summary["B4"].value, "=SUM(B2:B3)")
        self.assertEqual(len(summary._charts), 1)

    @unittest.skipUnless(_available("docx"), "python-docx is not installed")
    def test_docx_headings_paragraphs_and_source_note(self) -> None:
        import docx

        document = docx.Document(str(self.out / "backlog_report.docx"))
        styles = [paragraph.style.name for paragraph in document.paragraphs]
        self.assertIn("Heading 1", styles)
        self.assertIn("Heading 2", styles)
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("Backlog review", texts)
        self.assertTrue(any(text.startswith("Source: branchyard console fixture generator") for text in texts))

    @unittest.skipUnless(_available("pptx"), "python-pptx is not installed")
    def test_pptx_three_slides_with_native_chart(self) -> None:
        from pptx import Presentation

        presentation = Presentation(str(self.out / "backlog_briefing.pptx"))
        slides = list(presentation.slides)
        self.assertEqual(len(slides), 3)
        titles = []
        charts = 0
        for slide in slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    charts += 1
                    categories = list(shape.chart.plots[0].categories)
                    self.assertEqual(categories, ["Q1", "Q2"])
                known = {"Backlog overview", "Tickets closed by quarter", "Next actions"}
                if shape.has_text_frame and shape.text in known:
                    titles.append(shape.text)
        self.assertEqual(charts, 1)
        self.assertEqual(titles, ["Backlog overview", "Tickets closed by quarter", "Next actions"])

    def test_refuses_existing_directory(self) -> None:
        with self.assertRaises(branchyard_artifacts.ArtifactError):
            console_fixtures.generate(self.out)
        # The refused bundle is untouched.
        self.assertIn("manifest.json", {p.name for p in self.out.iterdir()})

    def test_refuses_symlink_destination(self) -> None:
        target = self.scratch / "real"
        target.mkdir()
        link = self.scratch / "link"
        link.symlink_to(target)
        with self.assertRaises(branchyard_artifacts.ArtifactError):
            console_fixtures.generate(link)
        self.assertEqual(list(target.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
