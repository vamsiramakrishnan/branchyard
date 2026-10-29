"""Actual source snapshots and citation-integrity failures, without network or models."""

import copy
import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "sdk/python"))

from branchyard_artifacts import evidence  # noqa: E402
from branchyard_artifacts.common import ArtifactError, digest  # noqa: E402

FORMATS = all(importlib.util.find_spec(name) for name in ["docx", "openpyxl", "pypdf", "reportlab"])


class EvidenceWorkspace(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "original.txt"
        self.source.write_bytes(b"Annual release: 552 hours.\nCapacity value: USD 41400.\n")
        self.capture_spec = {
            "schema_version": 1,
            "file": "original.txt",
            "source_id": "annual",
            "origin": {
                "uri": "https://example.invalid/report?period=2026",
                "retrieved_at": "2026-10-29T00:00:00Z",
                "title": "Annual capacity report",
            },
        }

    def spec(self, value, name="request.json"):
        path = self.root / name
        path.write_text(json.dumps(value))
        return path

    def captured(self):
        target = self.root / "evidence"
        result = evidence.capture(self.spec(self.capture_spec, "capture.json"), target)
        return target, result

    def claims(self, result):
        return {
            "schema_version": 1,
            "sources": [{"path": "evidence/index.json", "sha256": result["sha256"]}],
            "claims": [
                {
                    "id": "released-capacity",
                    "text": "The stated release is 552 hours.",
                    "evidence": [{"source_id": "annual", "locator": "line:1", "quote": "552 hours"}],
                }
            ],
        }


class EvidenceTests(EvidenceWorkspace):
    def test_capture_copies_exact_bytes_and_records_unverified_origin(self):
        target, result = self.captured()
        index = json.loads((target / "index.json").read_text())
        extracted = json.loads((target / "extraction.json").read_text())
        self.assertEqual((target / "source.txt").read_bytes(), self.source.read_bytes())
        self.assertEqual(index["source"]["sha256"], digest(self.source))
        self.assertEqual(index["source"]["size"], len(self.source.read_bytes()))
        self.assertEqual(index["source"]["origin"], self.capture_spec["origin"])
        self.assertIs(index["source"]["retrieval_verified"], False)
        self.assertIs(result["retrieval_verified"], False)
        self.assertEqual(result["extraction_review"], "not_performed")
        self.assertEqual(
            extracted["segments"][0], {"id": "annual/line:1", "locator": "line:1", "text": "Annual release: 552 hours."}
        )
        self.assertEqual(extracted["source_sha256"], digest(self.source))
        self.assertEqual(index["extraction"]["sha256"], digest(target / "extraction.json"))
        self.assertEqual({path.name for path in target.iterdir()}, {"source.txt", "index.json", "extraction.json"})

    def test_capture_is_deterministic_and_default_identity_is_content_bound(self):
        del self.capture_spec["source_id"]
        one = evidence.capture(self.spec(self.capture_spec), self.root / "one")
        two = evidence.capture(self.spec(self.capture_spec), self.root / "two")
        self.assertEqual(one["sha256"], two["sha256"])
        self.assertEqual(one["source_id"], "sha256:" + digest(self.source))
        self.assertEqual(
            (self.root / "one/extraction.json").read_bytes(), (self.root / "two/extraction.json").read_bytes()
        )

    def test_capture_failure_does_not_publish_partial_bundle_or_overwrite_existing(self):
        target = self.root / "evidence"
        original = Path.write_bytes

        def interrupted(path, data):
            if path.name == "extraction.json":
                raise OSError("interrupted staging")
            return original(path, data)

        spec = self.spec(self.capture_spec)
        with (
            mock.patch.object(Path, "write_bytes", interrupted),
            self.assertRaisesRegex(OSError, "interrupted staging"),
        ):
            evidence.capture(spec, target)
        self.assertFalse(target.exists())
        self.assertFalse(list(self.root.glob(".evidence-*")))
        evidence.capture(spec, target)
        before = (target / "index.json").read_bytes()
        with self.assertRaisesRegex(ArtifactError, "exists"):
            evidence.capture(spec, target)
        self.assertEqual((target / "index.json").read_bytes(), before)
        self.assertEqual(self.source.read_text().splitlines()[0], "Annual release: 552 hours.")

    def test_capture_refuses_symlinks_paths_special_files_and_unknown_suffixes(self):
        (self.root / "alias.txt").symlink_to(self.source)
        (self.root / "linked").symlink_to(self.root, target_is_directory=True)
        os.mkfifo(self.root / "pipe.txt")
        (self.root / "unknown.bin").write_bytes(b"data")
        for name in [
            "alias.txt",
            "linked/original.txt",
            "../original.txt",
            str(self.source),
            "pipe.txt",
            "unknown.bin",
            ".",
        ]:
            value = {**self.capture_spec, "file": name}
            with self.subTest(file=name), self.assertRaises(ArtifactError):
                evidence.capture(self.spec(value), self.root / "target")
        spec = self.spec(self.capture_spec)
        linked_spec = self.root / "linked-spec.json"
        linked_spec.symlink_to(spec)
        with self.assertRaises(OSError):
            evidence.capture(linked_spec, self.root / "target")
        os.mkfifo(self.root / "spec-pipe.json")
        with self.assertRaisesRegex(ArtifactError, "regular file"):
            evidence.capture(self.root / "spec-pipe.json", self.root / "target")

    def test_capture_validates_closed_requests_metadata_and_bounds(self):
        bad = [
            {**self.capture_spec, "fetch": True},
            {**self.capture_spec, "schema_version": True},
            {**self.capture_spec, "source_id": "../wrong"},
            {**self.capture_spec, "file": ""},
            {**self.capture_spec, "origin": {**self.capture_spec["origin"], "verified": True}},
            {
                **self.capture_spec,
                "origin": {**self.capture_spec["origin"], "uri": "https://user:secret@example.invalid/a"},
            },
            {**self.capture_spec, "origin": {**self.capture_spec["origin"], "uri": "relative/path"}},
            {**self.capture_spec, "origin": {**self.capture_spec["origin"], "retrieved_at": "2026-10-29T00:00:00"}},
            {**self.capture_spec, "origin": {**self.capture_spec["origin"], "title": ""}},
        ]
        for value in bad:
            with self.subTest(value=value), self.assertRaises(ValueError):
                evidence.capture(self.spec(value), self.root / "target")
        with mock.patch.object(evidence, "MAX_SOURCE_BYTES", 2), self.assertRaisesRegex(ArtifactError, "byte limit"):
            evidence.capture(self.spec(self.capture_spec), self.root / "target")
        with mock.patch.object(evidence, "MAX_TEXT_CHARS", 2), self.assertRaisesRegex(ArtifactError, "text limit"):
            evidence.capture(self.spec(self.capture_spec), self.root / "target")
        with mock.patch.object(evidence, "MAX_SEGMENTS", 1), self.assertRaisesRegex(ArtifactError, "segment"):
            evidence.capture(self.spec(self.capture_spec), self.root / "target")
        spec = self.spec([self.capture_spec])
        with self.assertRaisesRegex(ArtifactError, "JSON object"):
            evidence.capture(spec, self.root / "target")
        spec.write_text('{"schema_version":NaN}')
        with self.assertRaisesRegex(ArtifactError, "non-finite"):
            evidence.capture(spec, self.root / "target")

    def test_utf8_empty_sources_and_csv_preserve_identifiers_and_literal_na(self):
        self.source.write_text("Café — source\n\nSecond line")
        target, _ = self.captured()
        value = json.loads((target / "extraction.json").read_text())
        self.assertEqual([segment["text"] for segment in value["segments"]], ["Café — source", "", "Second line"])
        self.source.write_bytes(b"")
        result = evidence.capture(self.spec(self.capture_spec), self.root / "empty")
        self.assertEqual(result["segments"], 0)
        self.assertEqual(result["warnings"][0]["kind"], "no_extractable_text")
        csv_path = self.root / "source.csv"
        csv_path.write_text('id,note\n00012,NA\n00013,"two\nlines"\n')
        result = evidence.capture(self.spec({**self.capture_spec, "file": "source.csv"}), self.root / "csv")
        extracted = json.loads((self.root / "csv/extraction.json").read_text())
        segments = {segment["locator"]: segment["text"] for segment in extracted["segments"]}
        self.assertEqual(segments["row:2/column:1"], "00012")
        self.assertEqual(segments["row:2/column:2"], "NA")
        self.assertEqual(segments["row:3/column:2"], "two\nlines")
        self.assertEqual(result["segments"], 6)

    def test_quote_checks_do_not_claim_semantic_or_retrieval_verification(self):
        _, captured = self.captured()
        value = self.claims(captured)
        # A quote match is still not evidence that an unrelated claim follows.
        value["claims"][0]["text"] = "This investment is certain to succeed."
        out = self.root / "claim-report.json"
        result = evidence.verify_claims(self.spec(value), out)
        report = json.loads(out.read_text())
        self.assertEqual(result["status"], "passed")
        self.assertEqual(report["citation_integrity"], "passed")
        self.assertEqual(report["semantic_review"], "not_performed")
        self.assertEqual(report["claims"][0]["semantic_review"], "not_performed")
        self.assertIs(report["retrieval_verified"], False)
        self.assertEqual(report["sources"][0]["source_sha256"], digest(self.source))

    def test_quote_checks_reject_wrong_locator_unknown_source_and_nonexact_quote(self):
        _, captured = self.captured()
        for field, replacement, reason in [
            ("locator", "line:2", "quote_missing"),
            ("locator", "line:999", "unknown_locator"),
            ("quote", "552  hours", "quote_missing"),
            ("source_id", "unrelated", "unknown_source"),
        ]:
            value = self.claims(captured)
            value["claims"][0]["evidence"][0][field] = replacement
            out = self.root / "claim-report.json"
            with self.subTest(field=field, replacement=replacement):
                result = evidence.verify_claims(self.spec(value), out)
                self.assertEqual(result["status"], "failed")
                report = json.loads(out.read_text())
                self.assertEqual(report["claims"][0]["evidence"][0]["reason"], reason)

    def test_source_extraction_and_index_mutations_preserve_previous_report(self):
        target, captured = self.captured()
        value = self.claims(captured)
        out = self.root / "report.json"
        out.write_bytes(b"previous report")
        for path in [target / "source.txt", target / "extraction.json", target / "index.json"]:
            before = path.read_bytes()
            path.write_bytes(before + b"\nchanged")
            with self.subTest(path=path.name), self.assertRaisesRegex(ArtifactError, "SHA-256 mismatch"):
                evidence.verify_claims(self.spec(value), out)
            self.assertEqual(out.read_bytes(), b"previous report")
            path.write_bytes(before)

    def test_recomputed_receipt_hashes_cannot_authenticate_invented_extraction(self):
        target, captured = self.captured()
        extracted_path = target / "extraction.json"
        extraction = json.loads(extracted_path.read_text())
        extraction["segments"][0]["text"] = "Invented release: 999999 hours."
        extracted_path.write_text(json.dumps(extraction))
        index_path = target / "index.json"
        index = json.loads(index_path.read_text())
        index["extraction"]["sha256"] = digest(extracted_path)
        index_path.write_text(json.dumps(index))
        captured["sha256"] = digest(index_path)
        value = self.claims(captured)
        value["claims"][0]["evidence"][0]["quote"] = "999999 hours"
        with self.assertRaisesRegex(ArtifactError, "does not reproduce"):
            evidence.verify_claims(self.spec(value), self.root / "result.json")

    def test_quote_verifier_refuses_input_overwrites_and_capture_path_escapes(self):
        target, captured = self.captured()
        value = self.claims(captured)
        spec = self.spec(value)
        for out in [spec, target / "source.txt", target / "new-report.json"]:
            with self.subTest(out=out), self.assertRaisesRegex(ArtifactError, "preserve|outside captured"):
                evidence.verify_claims(spec, out)
        source_path = target / "source.txt"
        original = source_path.read_bytes()
        source_path.unlink()
        source_path.symlink_to(self.source)
        with self.assertRaisesRegex(ArtifactError, "Symlink"):
            evidence.verify_claims(spec, self.root / "result.json")
        source_path.unlink()
        source_path.write_bytes(original)
        value["sources"][0]["path"] = "../evidence/index.json"
        with self.assertRaisesRegex(ArtifactError, "relative"):
            evidence.verify_claims(self.spec(value), self.root / "result.json")

    def test_claim_request_is_closed_bounded_and_rejects_empty_evidence(self):
        _, captured = self.captured()
        base = self.claims(captured)
        bad = [
            {**base, "model": "forbidden"},
            {**base, "sources": []},
            {**base, "claims": []},
            {**base, "sources": [*base["sources"], *base["sources"]]},
            {**base, "claims": [*base["claims"], *base["claims"]]},
        ]
        for field, replacement in [("quote", ""), ("locator", ""), ("unrecognized", True)]:
            value = copy.deepcopy(base)
            value["claims"][0]["evidence"][0][field] = replacement
            bad.append(value)
        for replacement in [[], "wrong", [{}]]:
            value = copy.deepcopy(base)
            value["claims"][0]["evidence"] = replacement
            bad.append(value)
        for replacement in [None, "not-a-sha", captured["sha256"].upper()]:
            value = copy.deepcopy(base)
            value["sources"][0]["sha256"] = replacement
            bad.append(value)
        for value in bad:
            with self.subTest(value=value), self.assertRaises(ArtifactError):
                evidence.verify_claims(self.spec(value), self.root / "result.json")


@unittest.skipUnless(FORMATS, "install artifact libraries for real PDF/DOCX/XLSX fixtures")
class FormatEvidenceTests(EvidenceWorkspace):
    def test_actual_pdf_pages_and_blank_page_warning(self):
        from reportlab.pdfgen.canvas import Canvas

        file = self.root / "source.pdf"
        canvas = Canvas(str(file), invariant=1)
        canvas.drawString(72, 720, "Annual release: 552 hours.")
        canvas.showPage()
        canvas.showPage()
        canvas.save()
        value = {**self.capture_spec, "file": "source.pdf"}
        result = evidence.capture(self.spec(value), self.root / "evidence")
        extraction = json.loads((self.root / "evidence/extraction.json").read_text())
        self.assertEqual(result["segments"], 2)
        self.assertIn("552 hours", extraction["segments"][0]["text"])
        self.assertEqual(extraction["segments"][1]["locator"], "page:2")
        self.assertEqual(result["warnings"][0]["kind"], "no_extractable_text")
        self.assertIn("OCR did not run", result["warnings"][0]["next"])
        claims = self.claims(result)
        claims["claims"][0]["evidence"][0]["locator"] = "page:1"
        self.assertEqual(evidence.verify_claims(self.spec(claims), self.root / "report.json")["status"], "passed")
        with mock.patch.object(evidence, "MAX_PAGES", 1), self.assertRaisesRegex(ArtifactError, "page limit"):
            evidence.capture(self.spec(value), self.root / "too-many")

    def test_actual_docx_body_and_table_locators_with_explicit_scope(self):
        from docx import Document

        file = self.root / "source.docx"
        document = Document()
        document.add_paragraph("Annual release: 552 hours.")
        document.add_table(rows=1, cols=2).cell(0, 1).text = "USD 41400"
        document.sections[0].header.paragraphs[0].text = "Excluded header"
        document.save(file)
        result = evidence.capture(self.spec({**self.capture_spec, "file": file.name}), self.root / "evidence")
        extraction = json.loads((self.root / "evidence/extraction.json").read_text())
        self.assertEqual(extraction["extractor"]["options"]["scope"], "body_paragraphs_and_tables")
        self.assertEqual(extraction["segments"][0]["locator"], "paragraph:1")
        self.assertEqual(extraction["segments"][2]["locator"], "table:1/row:1/cell:2")
        self.assertEqual(extraction["segments"][2]["text"], "USD 41400")
        claims = self.claims(result)
        claims["claims"][0]["evidence"][0]["locator"] = "paragraph:1"
        self.assertEqual(evidence.verify_claims(self.spec(claims), self.root / "report.json")["status"], "passed")

    def test_actual_xlsx_cells_preserve_stored_formulas_strings_and_dates(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Capacity Plan"
        sheet.append(["00012", "NA", "=12*46", datetime(2026, 10, 29)])
        file = self.root / "source.xlsx"
        workbook.save(file)
        workbook.close()
        result = evidence.capture(self.spec({**self.capture_spec, "file": file.name}), self.root / "evidence")
        extraction = json.loads((self.root / "evidence/extraction.json").read_text())
        segments = {item["locator"]: item["text"] for item in extraction["segments"]}
        self.assertEqual(segments["sheet:Capacity Plan!A1"], "00012")
        self.assertEqual(segments["sheet:Capacity Plan!B1"], "NA")
        self.assertEqual(segments["sheet:Capacity Plan!C1"], "=12*46")
        self.assertEqual(segments["sheet:Capacity Plan!D1"], "2026-10-29T00:00:00")
        self.assertEqual(extraction["extractor"]["options"]["formulas"], "stored_text_not_evaluated")
        claims = self.claims(result)
        claims["claims"][0]["evidence"][0].update(locator="sheet:Capacity Plan!A1", quote="00012")
        self.assertEqual(evidence.verify_claims(self.spec(claims), self.root / "report.json")["status"], "passed")
        with mock.patch.object(evidence, "MAX_SEGMENTS", 2), self.assertRaisesRegex(ArtifactError, "cell limit"):
            evidence.capture(self.spec({**self.capture_spec, "file": file.name}), self.root / "too-many")

    def test_encrypted_pdf_and_expanded_office_bounds_are_refused(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("not-a-provider-credential")
        file = self.root / "encrypted.pdf"
        with file.open("wb") as target:
            writer.write(target)
        with self.assertRaisesRegex(ArtifactError, "Encrypted"):
            evidence.capture(self.spec({**self.capture_spec, "file": file.name}), self.root / "encrypted")
        fake = io.BytesIO()
        with evidence.zipfile.ZipFile(fake, "w") as archive:
            archive.writestr("fake", b"bytes")
        file = self.root / "large.docx"
        file.write_bytes(fake.getvalue())
        info = mock.Mock(file_size=129 * 1024 * 1024)
        with (
            mock.patch.object(evidence.zipfile.ZipFile, "infolist", return_value=[info]),
            self.assertRaisesRegex(ArtifactError, "expanded byte"),
        ):
            evidence.capture(self.spec({**self.capture_spec, "file": file.name}), self.root / "large")


if __name__ == "__main__":
    unittest.main()
