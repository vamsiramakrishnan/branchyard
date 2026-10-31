"""Bounded table, page, image and text preview tests. Hermetic; no model calls."""

from __future__ import annotations

import contextlib
import datetime
import hashlib
import importlib.util
import io
import json
import shutil
import stat
import subprocess
import sys
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path
from typing import Any
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
SDK = ROOT / "sdk/python"
sys.path.insert(0, str(SDK))

from branchyard_artifacts import (  # noqa: E402
    ArtifactError,
    Unavailable,
    docx,
    native,
    pdf,
    pptx,
    preview,
    xlsx,
)
from branchyard_artifacts.__main__ import main  # noqa: E402

EXAMPLES = SDK / "branchyard_artifacts/examples"
HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
HAS_DOCX = importlib.util.find_spec("docx") is not None
HAS_PPTX = importlib.util.find_spec("pptx") is not None
HAS_PDF = importlib.util.find_spec("pypdf") is not None and importlib.util.find_spec("reportlab") is not None
HAS_PIL = importlib.util.find_spec("PIL") is not None
HAS_FONTS = shutil.which("fc-match") is not None
HAS_LIBREOFFICE = (shutil.which("libreoffice") or shutil.which("soffice")) is not None
HAS_PDFTOPPM = shutil.which("pdftoppm") is not None


def _minimal_office_zip(fmt: str, extra_entries: dict[str, bytes] | None = None, content_type_extra: str = "") -> bytes:
    """Build a minimal structurally valid OOXML archive for archive-boundary tests."""
    part = {
        "docx": "word/document.xml",
        "xlsx": "xl/workbook.xml",
        "pptx": "ppt/presentation.xml",
    }[fmt]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            + content_type_extra
            + "</Types>",
        )
        archive.writestr(part, "<root/>")
        for name, data in (extra_entries or {}).items():
            archive.writestr(name, data)
    return buffer.getvalue()


def _set_first_sheet_dimension(path: Path, ref: str) -> None:
    """Rewrite the first worksheet's `<dimension>` the way a careless writer leaves it."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                start = data.index(b"<dimension ")
                end = data.index(b"/>", start) + 2
                data = data[:start] + b'<dimension ref="' + ref.encode() + b'"/>' + data[end:]
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())


class PreviewTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_cli(self, *args: object) -> tuple[int, dict]:
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            code = main([str(arg) for arg in args])
        return code, json.loads(stdout.getvalue())


class CliAndZeroDependencyTests(PreviewTestCase):
    def test_preview_parser_does_not_import_optional_libraries(self) -> None:
        program = """
import sys
sys.path.insert(0, sys.argv[1])
import branchyard, branchyard_client, branchyard_artifacts
from branchyard_artifacts.__main__ import parser
args = parser().parse_args([
    'preview', '--file', 'input.csv', '--out', 'out',
    '--view', 'auto', '--page', '1', '--sheet', 'Data', '--offset', '0', '--limit', '50'
])
assert args.format == 'preview'
assert args.view == 'auto'
assert args.page == 1
assert args.sheet == 'Data'
assert args.offset == 0
assert args.limit == 50
assert not any(
    name in sys.modules
    for name in ['docx', 'pptx', 'openpyxl', 'pypdf', 'reportlab', 'jsonschema', 'PIL', 'pandas', 'pint']
)
"""
        result = subprocess.run([sys.executable, "-I", "-c", program, str(SDK)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_cli_contract_exit_codes_for_passed_failed_and_unavailable(self) -> None:
        csv_path = self.root / "sample.csv"
        csv_path.write_text("id,value\n001,10\n", encoding="utf-8")
        out_dir = self.root / "out"
        code, payload = self.run_cli("preview", "--file", csv_path, "--out", out_dir)
        self.assertEqual(code, 0, payload)
        self.assertEqual(payload["status"], "passed")
        self.assertEqual(payload["kind"], "table")

        bad_code, bad_payload = self.run_cli("preview", "--file", csv_path, "--out", out_dir, "--offset", "-1")
        self.assertEqual(bad_code, 1, bad_payload)
        self.assertEqual(bad_payload["status"], "failed")
        self.assertEqual(bad_payload["error"]["kind"], "invalid_artifact")

        with mock.patch(
            "branchyard_artifacts.preview.library", side_effect=Unavailable("Missing openpyxl; install extras")
        ):
            xlsx_path = self.root / "sample.xlsx"
            xlsx_path.write_bytes(_minimal_office_zip("xlsx"))
            unavail_code, unavail_payload = self.run_cli("preview", "--file", xlsx_path, "--out", out_dir)
        self.assertEqual(unavail_code, 3, unavail_payload)
        self.assertEqual(unavail_payload["status"], "unavailable")
        self.assertEqual(unavail_payload["error"]["kind"], "unavailable")


class CsvPreviewTests(PreviewTestCase):
    def test_csv_preserves_leading_zeros_formulas_and_inert_markup_without_modifying_source(self) -> None:
        source = self.root / "evidence.csv"
        content = (
            "account_id,formula,markup,notes\n"
            "00042,=SUM(A1:A2),\"<script>alert('xss')</script>\",+1+2\n"
            '00007,"=CMD|\'/C calc\'!A0","<img src=x onerror=alert(1)>","@SUM(1,2)"\n'
        )
        source.write_text(content, encoding="utf-8")
        before = source.read_bytes()
        out_dir = self.root / "preview-out"

        code, result = self.run_cli(
            "preview",
            "--file",
            source,
            "--out",
            out_dir,
            "--view",
            "auto",
            "--page",
            "1",
            "--offset",
            "0",
            "--limit",
            "100",
        )
        self.assertEqual(code, 0, result)
        self.assertEqual(result["status"], "passed")
        self.assertEqual(result["kind"], "table")
        self.assertEqual(result["format"], "csv")
        self.assertEqual(result["columns"], ["account_id", "formula", "markup", "notes"])
        self.assertEqual(
            result["rows"],
            [
                ["00042", "=SUM(A1:A2)", "<script>alert('xss')</script>", "+1+2"],
                ["00007", "=CMD|'/C calc'!A0", "<img src=x onerror=alert(1)>", "@SUM(1,2)"],
            ],
        )
        self.assertEqual(result["sheets"], [])
        self.assertIsNone(result["sheet"])
        self.assertEqual(result["offset"], 0)
        self.assertIsNone(result["next_offset"])
        self.assertEqual(result["warnings"], [])
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse(out_dir.exists())

    def test_csv_pagination_and_truncation_bounds(self) -> None:
        source = self.root / "paged.csv"
        lines = ["c1,c2,c3,c4"]
        for idx in range(5):
            lines.append(f"row{idx},{'x' * 30},v{idx},extra{idx}")
        source.write_text("\n".join(lines) + "\n", encoding="utf-8")
        out_dir = self.root / "out"

        first = preview.preview(source, out_dir, view="table", offset=0, limit=2)
        self.assertEqual(len(first["rows"]), 2)
        self.assertEqual(first["rows"][0][0], "row0")
        self.assertEqual(first["rows"][1][0], "row1")
        self.assertEqual(first["offset"], 0)
        self.assertEqual(first["next_offset"], 2)

        second = preview.preview(source, out_dir, view="table", offset=2, limit=2)
        self.assertEqual(len(second["rows"]), 2)
        self.assertEqual(second["rows"][0][0], "row2")
        self.assertEqual(second["rows"][1][0], "row3")
        self.assertEqual(second["offset"], 2)
        self.assertEqual(second["next_offset"], 4)

        last = preview.preview(source, out_dir, view="table", offset=4, limit=2)
        self.assertEqual(len(last["rows"]), 1)
        self.assertEqual(last["rows"][0][0], "row4")
        self.assertEqual(last["offset"], 4)
        self.assertIsNone(last["next_offset"])

        with (
            mock.patch.object(preview, "MAX_COLUMNS", 2),
            mock.patch.object(preview, "MAX_CELL_CHARS", 10),
            mock.patch.object(preview, "MAX_SCAN_ROWS", 3),
        ):
            bounded = preview.preview(source, out_dir, view="table", offset=0, limit=3)
        self.assertEqual(bounded["columns"], ["c1", "c2"])
        self.assertEqual(len(bounded["rows"]), 3)
        self.assertEqual(bounded["rows"][0], ["row0", "x" * 10])
        self.assertIsNone(bounded["next_offset"])
        self.assertTrue(any("columns" in w for w in bounded["warnings"]))
        self.assertTrue(any("characters" in w for w in bounded["warnings"]))
        self.assertTrue(any("rows" in w for w in bounded["warnings"]))

    def test_csv_column_warning_requires_a_value_beyond_the_column_limit(self) -> None:
        padded = self.root / "padded.csv"
        padded.write_text("a,b,c" + "," * 150 + "\n1,2,3" + "," * 150 + "\n", encoding="utf-8")
        result = preview.preview(padded, self.root / "out")
        self.assertEqual(result["columns"][:4], ["a", "b", "c", "column_4"])
        self.assertEqual(len(result["columns"]), 100)
        self.assertEqual(result["rows"], [["1", "2", "3", *([""] * 97)]])
        self.assertEqual(result["warnings"], [])

        wide = self.root / "wide.csv"
        wide.write_text("a,b,c" + "," * 150 + "\n1,2,3" + "," * 97 + ",far\n", encoding="utf-8")
        result = preview.preview(wide, self.root / "out")
        self.assertEqual(len(result["columns"]), 100)
        self.assertEqual(result["warnings"], ["Truncated columns to 100"])

    def test_csv_ragged_rows_and_invalid_arguments_are_handled(self) -> None:
        source = self.root / "ragged.csv"
        source.write_text(",b\n1\n1,2,3\n", encoding="utf-8")
        result = preview.preview(source, self.root / "out")
        self.assertEqual(result["columns"], ["column_1", "b", "column_3"])
        self.assertEqual(result["rows"], [["1", "", ""], ["1", "2", "3"]])

        invalid_calls: list[dict[str, Any]] = [
            {"view": "pages"},
            {"view": "invalid"},
            {"sheet": "Sheet1"},
            {"page": 2},
            {"page": 0},
            {"page": True},
            {"offset": -1},
            {"offset": True},
            {"limit": 0},
            {"limit": 501},
            {"limit": False},
        ]
        for kwargs in invalid_calls:
            with self.subTest(kwargs=kwargs), self.assertRaises(ArtifactError):
                preview.preview(source, self.root / "out", **kwargs)

        for bad_bytes in [b"", b"\x00binary", b"\xff\xfeinvalid-utf8", b",\n", b'a,b\n"unterminated\n']:
            bad_csv = self.root / "bad.csv"
            bad_csv.write_bytes(bad_bytes)
            with self.subTest(bad_bytes=bad_bytes), self.assertRaises(ArtifactError):
                preview.preview(bad_csv, self.root / "out")


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl is required for real XLSX fixture tests")
class XlsxPreviewTests(PreviewTestCase):
    def make_workbook(self) -> Path:
        import openpyxl

        path = self.root / "workbook.xlsx"
        wb = openpyxl.Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["Metric", "Value", "Notes"])
        summary.append(["Account", "00012", "<script>inert</script>"])
        summary.append(["Total", "=SUM(Data!B2:B4)", datetime.date(2026, 10, 29)])
        summary.append(["Active", True, Decimal("12.5")])

        data = wb.create_sheet("Data")
        data.append(["Quarter", "Revenue"])
        data.append(["Q1", 12])
        data.append(["Q2", 14])
        data.append(["Q3", 18])
        data.append(["Total", "=SUM(B2:B4)"])

        empty = wb.create_sheet("Empty")
        del empty
        wb.save(path)
        return path

    def test_xlsx_sheet_selection_formulas_as_literal_text_and_pagination(self) -> None:
        source = self.make_workbook()
        before = source.read_bytes()
        out_dir = self.root / "out"

        default_res = preview.preview(source, out_dir)
        self.assertEqual(default_res["status"], "passed")
        self.assertEqual(default_res["kind"], "table")
        self.assertEqual(default_res["format"], "xlsx")
        self.assertEqual(default_res["sheets"], ["Summary", "Data", "Empty"])
        self.assertEqual(default_res["sheet"], "Summary")
        self.assertEqual(default_res["columns"], ["Metric", "Value", "Notes"])
        self.assertEqual(
            default_res["rows"],
            [
                ["Account", "00012", "<script>inert</script>"],
                ["Total", "=SUM(Data!B2:B4)", "2026-10-29T00:00:00"],
                ["Active", True, 12.5],
            ],
        )
        self.assertTrue(any("formula" in w.lower() for w in default_res["warnings"]))

        code, page1 = self.run_cli(
            "preview",
            "--file",
            source,
            "--out",
            out_dir,
            "--view",
            "table",
            "--sheet",
            "Data",
            "--offset",
            "0",
            "--limit",
            "2",
        )
        self.assertEqual(code, 0, page1)
        self.assertEqual(page1["sheet"], "Data")
        self.assertEqual(page1["columns"], ["Quarter", "Revenue"])
        self.assertEqual(page1["rows"], [["Q1", 12], ["Q2", 14]])
        self.assertEqual(page1["next_offset"], 2)

        page2 = preview.preview(source, out_dir, view="table", sheet="Data", offset=2, limit=2)
        self.assertEqual(page2["rows"], [["Q3", 18], ["Total", "=SUM(B2:B4)"]])
        self.assertIsNone(page2["next_offset"])
        self.assertTrue(any("formula" in w.lower() for w in page2["warnings"]))

        empty_res = preview.preview(source, out_dir, view="table", sheet="Empty")
        self.assertEqual(empty_res["columns"], [])
        self.assertEqual(empty_res["rows"], [])
        self.assertIsNone(empty_res["next_offset"])

        with self.assertRaisesRegex(ArtifactError, "Worksheet does not exist"):
            preview.preview(source, out_dir, sheet="NonExistent")
        with self.assertRaisesRegex(ArtifactError, "sheet"):
            preview.preview(source, out_dir, sheet="   ")
        self.assertEqual(source.read_bytes(), before)

    def test_xlsx_column_truncation_and_formula_object_handling(self) -> None:
        source = self.make_workbook()
        out_dir = self.root / "out"
        with (
            mock.patch.object(preview, "MAX_COLUMNS", 2),
            mock.patch.object(preview, "MAX_CELL_CHARS", 8),
            mock.patch.object(preview, "MAX_SCAN_ROWS", 2),
        ):
            res = preview.preview(source, out_dir, sheet="Summary", offset=0, limit=2)
        self.assertEqual(res["columns"], ["Metric", "Value"])
        self.assertEqual(res["rows"], [["Account", "00012"], ["Total", "=SUM(Dat"]])
        self.assertIsNone(res["next_offset"])
        self.assertTrue(any("columns" in w for w in res["warnings"]))
        self.assertTrue(any("characters" in w for w in res["warnings"]))
        self.assertTrue(any("rows" in w for w in res["warnings"]))

        with mock.patch.object(preview, "MAX_COLUMNS", 3):
            exact = preview.preview(source, out_dir, sheet="Summary")
        self.assertEqual(exact["columns"], ["Metric", "Value", "Notes"])
        self.assertFalse(any("columns" in w for w in exact["warnings"]), exact["warnings"])

        class FakeArrayFormula:
            text = "=SUM(A1:A10)"

        warnings: list[str] = []
        self.assertEqual(preview._xlsx_cell(FakeArrayFormula(), warnings), "=SUM(A1:A10)")
        self.assertTrue(any("formula" in w.lower() for w in warnings))

    def test_xlsx_column_warning_requires_a_value_beyond_the_column_limit(self) -> None:
        import openpyxl
        from openpyxl.styles import Font

        def styled_empty_cell(sheet: Any) -> None:
            sheet.cell(row=1, column=150).font = Font(bold=True)

        def value_past_empty_column_101(sheet: Any) -> None:
            sheet.cell(row=2, column=150, value="note")

        cases: list[tuple[str, Any, list[str]]] = [
            ("narrow", lambda sheet: None, []),
            ("styled empty cell in column 150", styled_empty_cell, []),
            ("value in column 150 only", value_past_empty_column_101, ["Truncated columns to 100"]),
        ]
        for label, extend, expected in cases:
            with self.subTest(label):
                path = self.root / "narrow.xlsx"
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append(["Region", "Q1", "Q2"])
                sheet.append(["North", 1, 2])
                extend(sheet)
                workbook.save(path)
                result = preview.preview(path, self.root / "out")
                self.assertEqual(result["columns"], ["Region", "Q1", "Q2"])
                self.assertEqual(result["rows"], [["North", 1, 2]])
                self.assertEqual(result["warnings"], expected)

    def test_xlsx_value_in_column_101_warns_once_and_shows_100_columns(self) -> None:
        import openpyxl

        path = self.root / "wide.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append([f"c{index}" for index in range(1, 102)])
        sheet.append(list(range(1, 102)))
        sheet.append(list(range(1, 102)))
        workbook.save(path)
        result = preview.preview(path, self.root / "out")
        self.assertEqual(result["columns"], [f"c{index}" for index in range(1, 101)])
        self.assertEqual(result["rows"], [list(range(1, 101))] * 2)
        self.assertEqual(result["warnings"], ["Truncated columns to 100"])

    def test_xlsx_preview_reads_past_an_understated_dimension(self) -> None:
        import openpyxl

        path = self.root / "understated.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in (["Region", "Q1", "Q2"], ["North", 1, 2], ["South", 3, 4], ["East", 5, 6]):
            sheet.append(row)
        sheet.cell(row=4, column=150, value="note")
        workbook.save(path)
        _set_first_sheet_dimension(path, "A1:C2")
        result = preview.preview(path, self.root / "out")
        self.assertEqual(result["columns"], ["Region", "Q1", "Q2"])
        self.assertEqual(result["rows"], [["North", 1, 2], ["South", 3, 4], ["East", 5, 6]])
        self.assertEqual(result["warnings"], ["Truncated columns to 100"])


@unittest.skipUnless(HAS_PDF and HAS_PIL, "pypdf, reportlab and Pillow are required for page preview tests")
class SyntheticPageAndImagePreviewTests(PreviewTestCase):
    def make_pdf(self, name: str = "sample.pdf", pages: int = 3) -> Path:
        from reportlab.pdfgen.canvas import Canvas

        path = self.root / name
        canvas = Canvas(str(path), pagesize=(612, 792))
        for index in range(pages):
            canvas.drawString(72, 720, f"Synthetic page {index + 1}")
            canvas.showPage()
        canvas.save()
        return path

    def install_fake_converters(self, pdf_pages: int = 3, png_size: tuple[int, int] = (300, 400)) -> tuple[Path, Path]:
        """Create deterministic fake libreoffice and pdftoppm binaries on PATH."""
        bin_dir = self.root / "fake-bin"
        bin_dir.mkdir(exist_ok=True)
        log_file = self.root / "converter-calls.jsonl"
        template_pdf = self.make_pdf("template.pdf", pages=pdf_pages)
        template_png = self.root / "template.png"
        from PIL import Image

        Image.new("RGB", png_size, "navy").save(template_png, format="PNG")

        office_script = bin_dir / "libreoffice"
        office_script.write_text(
            f"""#!{sys.executable}
import json, pathlib, shutil, sys
args = sys.argv[1:]
outdir = pathlib.Path(args[args.index("--outdir") + 1])
source = pathlib.Path(args[-1])
profile_arg = next(a for a in args if a.startswith("-env:UserInstallation="))
profile_uri = profile_arg.split("=", 1)[1]
profile_path = pathlib.Path(profile_uri.removeprefix("file://"))
xcu = (profile_path / "user/registrymodifications.xcu").read_text(encoding="utf-8")
with open({str(log_file)!r}, "a", encoding="utf-8") as handle:
    handle.write(json.dumps({{"tool": "libreoffice", "args": args, "xcu": xcu}}) + "\\n")
shutil.copyfile({str(template_pdf)!r}, outdir / (source.stem + ".pdf"))
""",
            encoding="utf-8",
        )
        office_script.chmod(0o755)

        pdftoppm_script = bin_dir / "pdftoppm"
        pdftoppm_script.write_text(
            f"""#!{sys.executable}
import json, pathlib, shutil, sys
args = sys.argv[1:]
first_page = int(args[args.index("-f") + 1])
last_page = int(args[args.index("-l") + 1])
prefix = pathlib.Path(args[-1])
with open({str(log_file)!r}, "a", encoding="utf-8") as handle:
    handle.write(json.dumps({{"tool": "pdftoppm", "args": args, "first": first_page, "last": last_page}}) + "\\n")
assert first_page == last_page, "pdftoppm must rasterize only the single requested page"
shutil.copyfile({str(template_png)!r}, prefix.parent / f"{{prefix.name}}-{{first_page}}.png")
""",
            encoding="utf-8",
        )
        pdftoppm_script.chmod(0o755)
        return bin_dir, log_file

    def test_docx_pptx_xlsx_and_pdf_single_page_selection_via_fake_converters(self) -> None:
        bin_dir, log_file = self.install_fake_converters(pdf_pages=3, png_size=(320, 240))
        docx_file = self.root / "report.docx"
        pptx_file = self.root / "slides.pptx"
        xlsx_file = self.root / "sheet.xlsx"
        pdf_file = self.make_pdf("doc.pdf", pages=3)
        docx_file.write_bytes(_minimal_office_zip("docx"))
        pptx_file.write_bytes(_minimal_office_zip("pptx"))
        xlsx_file.write_bytes(_minimal_office_zip("xlsx"))

        def fake_which(name: str) -> str | None:
            candidate = bin_dir / name
            if candidate.is_file():
                return str(candidate)
            return None

        with mock.patch("shutil.which", side_effect=fake_which):
            for index, (source, expected_fmt, view) in enumerate(
                [
                    (docx_file, "docx", "auto"),
                    (pptx_file, "pptx", "auto"),
                    (xlsx_file, "xlsx", "pages"),
                    (pdf_file, "pdf", "auto"),
                ],
                1,
            ):
                before = source.read_bytes()
                out_dir = self.root / f"out-{index}"
                code, result = self.run_cli(
                    "preview",
                    "--file",
                    source,
                    "--out",
                    out_dir,
                    "--view",
                    view,
                    "--page",
                    "2",
                )
                self.assertEqual(code, 0, result)
                self.assertEqual(
                    result,
                    {
                        "status": "passed",
                        "kind": "pages",
                        "format": expected_fmt,
                        "page": 2,
                        "pages": 3,
                        "image": "page-2.png",
                        "width": 320,
                        "height": 240,
                        "warnings": [],
                    },
                )
                self.assertEqual(source.read_bytes(), before)
                self.assertEqual([p.name for p in out_dir.iterdir()], ["page-2.png"])

            with self.assertRaisesRegex(ArtifactError, "exceeds document page count"):
                preview.preview(pdf_file, self.root / "out-oob", page=4)
            self.assertFalse((self.root / "out-oob").exists())

        calls = [json.loads(line) for line in log_file.read_text(encoding="utf-8").splitlines()]
        office_calls = [c for c in calls if c["tool"] == "libreoffice"]
        raster_calls = [c for c in calls if c["tool"] == "pdftoppm"]
        self.assertEqual(len(office_calls), 3)
        for call in office_calls:
            self.assertIn("MacroSecurityLevel", call["xcu"])
            self.assertIn("<value>3</value>", call["xcu"])
            self.assertIn("BlockUntrustedRefererLinks", call["xcu"])
        self.assertEqual(len(raster_calls), 4)
        for call in raster_calls:
            self.assertEqual(call["first"], 2)
            self.assertEqual(call["last"], 2)

    def test_huge_pdf_page_is_bounded_before_rasterization(self) -> None:
        from pypdf import PdfWriter

        bin_dir, log_file = self.install_fake_converters(pdf_pages=1)
        source = self.root / "huge.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=1_000_000, height=1_000_000)
        with source.open("wb") as stream:
            writer.write(stream)

        def fake_which(name: str) -> str | None:
            candidate = bin_dir / name
            return str(candidate) if candidate.is_file() else None

        with mock.patch("shutil.which", side_effect=fake_which):
            preview.preview(source, self.root / "huge-preview")
        raster = next(json.loads(line) for line in log_file.read_text().splitlines() if '"pdftoppm"' in line)
        self.assertIn("-scale-to", raster["args"])
        scale = int(raster["args"][raster["args"].index("-scale-to") + 1])
        self.assertGreater(scale, 0)
        self.assertLessEqual(scale, preview.MAX_PREVIEW_DIMENSION)

    def test_oversized_rendered_page_is_scaled_and_converter_failures_leave_no_partial_outputs(self) -> None:
        bin_dir, _ = self.install_fake_converters(pdf_pages=2, png_size=(1200, 900))
        pdf_file = self.make_pdf("scale.pdf", pages=2)

        def fake_which(name: str) -> str | None:
            candidate = bin_dir / name
            return str(candidate) if candidate.is_file() else None

        out_dir = self.root / "scaled-out"
        with (
            mock.patch("shutil.which", side_effect=fake_which),
            mock.patch.object(preview, "MAX_PREVIEW_DIMENSION", 600),
        ):
            res = preview.preview(pdf_file, out_dir, page=1)
        self.assertEqual(res["width"], 600)
        self.assertEqual(res["height"], 450)
        self.assertTrue(any("600px" in w for w in res["warnings"]))

        failed_out = self.root / "failed-out"
        with (
            mock.patch("shutil.which", side_effect=fake_which),
            mock.patch("subprocess.run", side_effect=subprocess.TimeoutExpired(cmd="pdftoppm", timeout=1)),
            self.assertRaisesRegex(ArtifactError, "timed out"),
        ):
            preview.preview(pdf_file, failed_out, page=1)
        self.assertFalse(failed_out.exists())

    def test_image_preview_converts_to_safe_png_bounds_dimensions_and_rejects_svg_or_html(self) -> None:
        from PIL import Image

        jpg_path = self.root / "photo.jpg"
        Image.new("RGB", (800, 600), "crimson").save(jpg_path, format="JPEG")
        before = jpg_path.read_bytes()
        out_dir = self.root / "img-out"

        with mock.patch.object(preview, "MAX_PREVIEW_DIMENSION", 400):
            res = preview.preview(jpg_path, out_dir)
        self.assertEqual(
            res,
            {
                "status": "passed",
                "kind": "pages",
                "format": "jpeg",
                "page": 1,
                "pages": 1,
                "image": "page-1.png",
                "width": 400,
                "height": 300,
                "warnings": ["Scaled image within 400px dimension bound"],
            },
        )
        self.assertEqual(jpg_path.read_bytes(), before)
        rendered_png = out_dir / "page-1.png"
        self.assertTrue(rendered_png.is_file())
        with Image.open(rendered_png) as img:
            self.assertEqual(img.format, "PNG")
            self.assertEqual(img.size, (400, 300))

        with self.assertRaisesRegex(ArtifactError, "1 page"):
            preview.preview(jpg_path, self.root / "bad-page", page=2)

        with (
            mock.patch.object(preview, "MAX_IMAGE_DIMENSION", 100),
            self.assertRaisesRegex(ArtifactError, "dimensions exceed"),
        ):
            preview.preview(jpg_path, self.root / "too-big")
        self.assertFalse((self.root / "too-big").exists())

        for name, content in [
            ("vector.svg", b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'),
            ("page.html", b"<!DOCTYPE html><html><body><script>alert(1)</script></body></html>"),
            ("disguised.png", b'<svg xmlns="http://www.w3.org/2000/svg"></svg>'),
            ("corrupt.png", b"not-a-valid-png-payload"),
        ]:
            bad_path = self.root / name
            bad_path.write_bytes(content)
            with self.subTest(name=name), self.assertRaises(ArtifactError):
                preview.preview(bad_path, self.root / "rejected-out")
            self.assertFalse((self.root / "rejected-out").exists())


class TextAndSecurityBoundaryTests(PreviewTestCase):
    def test_text_markdown_and_json_previews_enforce_bounds(self) -> None:
        md_path = self.root / "notes.md"
        md_path.write_text("# Heading\nLiteral text.\n", encoding="utf-8")
        res = preview.preview(md_path, self.root / "out")
        self.assertEqual(res["status"], "passed")
        self.assertEqual(res["kind"], "text")
        self.assertEqual(res["format"], "markdown")
        self.assertFalse(res["truncated"])
        self.assertEqual(res["text"], "# Heading\nLiteral text.\n")

        with mock.patch.object(preview, "MAX_TEXT_CHARS", 9):
            trunc = preview.preview(md_path, self.root / "out")
        self.assertTrue(trunc["truncated"])
        self.assertEqual(trunc["text"], "# Heading")
        self.assertTrue(trunc["warnings"])

        bad_json = self.root / "broken.json"
        bad_json.write_text("{not valid json", encoding="utf-8")
        with self.assertRaisesRegex(ArtifactError, "Malformed JSON"):
            preview.preview(bad_json, self.root / "out")

    def test_rejects_symlinks_macros_malformed_archives_oversized_files_and_source_overwrite(self) -> None:
        csv_path = self.root / "valid.csv"
        csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
        original = csv_path.read_bytes()

        # Symlink input file
        link_file = self.root / "link.csv"
        link_file.symlink_to(csv_path)
        with self.assertRaisesRegex(ArtifactError, "symbolic link"):
            preview.preview(link_file, self.root / "out")

        # Symlink ancestor directory
        real_dir = self.root / "real"
        real_dir.mkdir()
        nested = real_dir / "nested.csv"
        nested.write_text("a,b\n1,2\n", encoding="utf-8")
        link_dir = self.root / "link-dir"
        link_dir.symlink_to(real_dir, target_is_directory=True)
        with self.assertRaisesRegex(ArtifactError, "symbolic link"):
            preview.preview(link_dir / "nested.csv", self.root / "out")

        # Symlink output directory
        out_target = self.root / "out-target"
        out_target.mkdir()
        link_out = self.root / "link-out"
        link_out.symlink_to(out_target, target_is_directory=True)
        with self.assertRaisesRegex(ArtifactError, "symbolic link"):
            preview.preview(csv_path, link_out)

        # Output equals source or is an existing file
        with self.assertRaisesRegex(ArtifactError, "overwrite|directory"):
            preview.preview(csv_path, csv_path)
        existing_file = self.root / "existing-file"
        existing_file.write_text("keep", encoding="utf-8")
        with self.assertRaisesRegex(ArtifactError, "directory"):
            preview.preview(csv_path, existing_file)
        self.assertEqual(csv_path.read_bytes(), original)
        self.assertEqual(existing_file.read_text(encoding="utf-8"), "keep")

        # Oversized input
        with (
            mock.patch.object(preview, "MAX_INPUT_BYTES", 4),
            self.assertRaisesRegex(ArtifactError, "byte limit"),
        ):
            preview.preview(csv_path, self.root / "out")

        # Macro-enabled extensions
        for ext in [".xlsm", ".docm", ".pptm", ".xlsb"]:
            macro_file = self.root / ("macro" + ext)
            macro_file.write_bytes(b"macro-content")
            with self.subTest(ext=ext), self.assertRaisesRegex(ArtifactError, "Macro-enabled"):
                preview.preview(macro_file, self.root / "out")

        # Embedded vbaProject.bin or macro content type inside .xlsx/.docx
        disguised_vba = self.root / "disguised.xlsx"
        disguised_vba.write_bytes(_minimal_office_zip("xlsx", {"xl/vbaProject.bin": b"VBA"}))
        with self.assertRaisesRegex(ArtifactError, "Macro-enabled"):
            preview.preview(disguised_vba, self.root / "out")

        disguised_ct = self.root / "disguised-ct.docx"
        disguised_ct.write_bytes(
            _minimal_office_zip(
                "docx",
                content_type_extra='<Override PartName="/word/vbaProject.bin" '
                'ContentType="application/vnd.ms-office.vbaProject"/>',
            )
        )
        with self.assertRaisesRegex(ArtifactError, "Macro-enabled"):
            preview.preview(disguised_ct, self.root / "out")

        # Traversal and symlink inside Office ZIP
        traversal_zip = self.root / "traversal.xlsx"
        traversal_zip.write_bytes(_minimal_office_zip("xlsx", {"../evil.xml": b"bad"}))
        with self.assertRaisesRegex(ArtifactError, "unsafe member path"):
            preview.preview(traversal_zip, self.root / "out")

        symlink_buf = io.BytesIO(_minimal_office_zip("xlsx"))
        with zipfile.ZipFile(symlink_buf, "a") as zf:
            info = zipfile.ZipInfo("xl/link.xml")
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
            zf.writestr(info, "/etc/passwd")
        symlink_zip = self.root / "symlink.xlsx"
        symlink_zip.write_bytes(symlink_buf.getvalue())
        with self.assertRaisesRegex(ArtifactError, "symbolic link member"):
            preview.preview(symlink_zip, self.root / "out")

        # Archive expansion / entry count bounds
        valid_xlsx = self.root / "valid-structure.xlsx"
        valid_xlsx.write_bytes(_minimal_office_zip("xlsx", {"xl/extra.xml": b"x" * 200}))
        with (
            mock.patch.object(preview, "MAX_ARCHIVE_ENTRIES", 1),
            self.assertRaisesRegex(ArtifactError, "entry count"),
        ):
            preview.preview(valid_xlsx, self.root / "out")
        with (
            mock.patch.object(preview, "MAX_ARCHIVE_EXPANDED_BYTES", 50),
            self.assertRaisesRegex(ArtifactError, "expanded|Expanded"),
        ):
            preview.preview(valid_xlsx, self.root / "out")

        # Missing required OOXML part
        wrong_part = self.root / "wrong.docx"
        wrong_part.write_bytes(_minimal_office_zip("xlsx"))
        with self.assertRaisesRegex(ArtifactError, "missing required document parts"):
            preview.preview(wrong_part, self.root / "out")

    def test_unavailable_tools_report_setup_instructions_without_partial_outputs(self) -> None:
        docx_file = self.root / "doc.docx"
        docx_file.write_bytes(_minimal_office_zip("docx"))
        out_dir = self.root / "unavail-out"
        with mock.patch("shutil.which", return_value=None):
            code, payload = self.run_cli("preview", "--file", docx_file, "--out", out_dir)
        self.assertEqual(code, 3, payload)
        self.assertEqual(payload["status"], "unavailable")
        self.assertIn("pdftoppm", payload["error"]["message"])
        self.assertIn("LibreOffice", payload["error"]["message"])
        self.assertFalse(out_dir.exists())


@unittest.skipUnless(
    HAS_DOCX and HAS_OPENPYXL and HAS_PPTX and HAS_PDF and HAS_PIL and HAS_FONTS and HAS_LIBREOFFICE and HAS_PDFTOPPM,
    "real Office/PDF libraries, LibreOffice and Poppler are required for live native render fixture tests",
)
class RealNativeRenderPreviewTests(PreviewTestCase):
    def test_real_libreoffice_and_poppler_render_single_page_previews_for_docx_xlsx_pptx_and_pdf(self) -> None:
        from PIL import Image

        docx_path = self.root / "real.docx"
        xlsx_path = self.root / "real.xlsx"
        pptx_path = self.root / "real.pptx"
        pdf_path = self.root / "real.pdf"
        docx.create(EXAMPLES / "docx.json", docx_path)
        xlsx.create(EXAMPLES / "xlsx.json", xlsx_path)
        pptx.create(EXAMPLES / "deck.json", pptx_path)
        pdf.create(EXAMPLES / "pdf.json", pdf_path)

        for fmt, source, view in [
            ("docx", docx_path, "auto"),
            ("xlsx", xlsx_path, "pages"),
            ("pptx", pptx_path, "auto"),
            ("pdf", pdf_path, "auto"),
        ]:
            with self.subTest(format=fmt):
                before_digest = hashlib.sha256(source.read_bytes()).hexdigest()
                out_dir = self.root / f"real-{fmt}-preview"
                result = preview.preview(source, out_dir, view=view, page=1)
                self.assertEqual(result["status"], "passed")
                self.assertEqual(result["kind"], "pages")
                self.assertEqual(result["format"], fmt)
                self.assertEqual(result["page"], 1)
                self.assertGreaterEqual(result["pages"], 1)
                self.assertEqual(result["image"], "page-1.png")
                self.assertGreater(result["width"], 0)
                self.assertGreater(result["height"], 0)
                image_path = out_dir / result["image"]
                self.assertTrue(image_path.is_file())
                self.assertEqual([item.name for item in out_dir.iterdir()], ["page-1.png"])
                with Image.open(image_path) as decoded:
                    self.assertEqual(decoded.size, (result["width"], result["height"]))
                self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before_digest)
                self.assertIsNotNone(native)


if __name__ == "__main__":
    unittest.main()
