"""Drive the shipped knowledge CLI with real local artifacts and no model or network."""

import contextlib
import importlib.util
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
SDK = ROOT / "sdk/python"
sys.path.insert(0, str(SDK))

from branchyard_artifacts.__main__ import main  # noqa: E402
from branchyard_artifacts.common import Unavailable  # noqa: E402


class KnowledgeCliTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def write(self, name, value):
        path = self.root / name
        path.write_text(json.dumps(value))
        return path

    def run_cli(self, *args):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            code = main(list(map(str, args)))
        return code, json.loads(output.getvalue())

    def test_parser_does_not_import_optional_libraries(self):
        program = """
import sys
sys.path.insert(0, sys.argv[1])
from branchyard_artifacts.__main__ import parser
for action in ['source-capture','table-extract','data-clean','calculate','claim-verify','inspect-layout']:
    parser().parse_args(['knowledge',action,'input','--out','output'])
assert not any(n in sys.modules for n in ['pandas','pint','pdfplumber','docx','openpyxl','pypdf'])
"""
        result = subprocess.run([sys.executable, "-I", "-c", program, str(SDK)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_source_claim_lifecycle_and_failed_evidence_exit_status(self):
        (self.root / "report.txt").write_text("552 hours of released capacity\n")
        capture = self.write(
            "capture.json",
            {
                "schema_version": 1,
                "file": "report.txt",
                "source_id": "capacity",
                "origin": {
                    "uri": "https://example.org/report",
                    "retrieved_at": "2026-10-29T00:00:00Z",
                    "title": "Illustrative fixture",
                },
            },
        )
        code, source = self.run_cli("knowledge", "source-capture", capture, "--out", self.root / "evidence")
        self.assertEqual(code, 0, source)
        claim = {
            "schema_version": 1,
            "sources": [{"path": "evidence/index.json", "sha256": source["sha256"]}],
            "claims": [
                {
                    "id": "capacity",
                    "text": "552 hours are released",
                    "evidence": [{"source_id": "capacity", "locator": "line:1", "quote": "552 hours"}],
                }
            ],
        }
        spec = self.write("claims.json", claim)
        target = self.root / "claim-report.json"
        code, value = self.run_cli("knowledge", "claim-verify", spec, "--out", target)
        self.assertEqual(code, 0, value)
        self.assertEqual(json.loads(target.read_text())["semantic_review"], "not_performed")
        claim["claims"][0]["evidence"][0]["quote"] = "999 hours"
        self.write("claims.json", claim)
        code, value = self.run_cli("knowledge", "claim-verify", spec, "--out", target)
        self.assertEqual(code, 1, value)
        self.assertEqual(json.loads(target.read_text())["claims"][0]["evidence"][0]["reason"], "quote_missing")

    @unittest.skipUnless(
        importlib.util.find_spec("pandas") and importlib.util.find_spec("pint"), "optional analysis libs"
    )
    def test_actual_cleaning_and_quantity_calculation_dispatch(self):
        (self.root / "sales.csv").write_text("account_id,region,amount\n00012, NA ,3.5\n00012, NA ,3.5\n")
        clean = self.write(
            "clean.json",
            {
                "schema_version": 1,
                "input": {"path": "sales.csv"},
                "trim": ["region"],
                "deduplicate": {"columns": ["account_id"], "keep": "first"},
            },
        )
        out = self.root / "cleaned"
        code, value = self.run_cli("knowledge", "data-clean", clean, "--out", out)
        self.assertEqual(code, 0, value)
        self.assertEqual((out / "data.csv").read_text(), "account_id,region,amount\n00012,NA,3.5\n")
        calculation = self.write(
            "calculation.json",
            {
                "schema_version": 1,
                "variables": {"distance": {"value": 120, "unit": "km"}},
                "expression": "distance / 2",
                "unit": "meter",
            },
        )
        target = self.root / "result.json"
        code, value = self.run_cli("knowledge", "calculate", calculation, "--out", target)
        self.assertEqual(code, 0, value)
        self.assertEqual(json.loads(target.read_text())["result"]["value"], 60000)

    def test_missing_capability_is_unavailable_without_implicit_install(self):
        spec = self.write(
            "calc.json", {"schema_version": 1, "variables": {}, "expression": "1", "unit": "dimensionless"}
        )
        with mock.patch("branchyard_artifacts.analysis.library", side_effect=Unavailable("Missing Pint")):
            code, value = self.run_cli("knowledge", "calculate", spec, "--out", self.root / "out.json")
        self.assertEqual(code, 3, value)
        self.assertFalse((self.root / "out.json").exists())

    @unittest.skipUnless(importlib.util.find_spec("docx"), "optional DOCX library")
    def test_actual_table_dispatch_preserves_location_and_original_bytes(self):
        import docx

        path = self.root / "table.docx"
        document = docx.Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "ID", "Hours"
        table.cell(1, 0).text, table.cell(1, 1).text = "00012", "552"
        document.save(path)
        before = path.read_bytes()
        code, value = self.run_cli("knowledge", "table-extract", path, "--out", self.root / "table.json")
        self.assertEqual(code, 0, value)
        self.assertEqual(value["tables"][0]["rows"][1], ["00012", "552"])
        self.assertEqual(path.read_bytes(), before)

    @unittest.skipUnless(
        shutil.which("pdftoppm") and all(importlib.util.find_spec(n) for n in ["reportlab", "pypdf", "PIL"]),
        "native PDF rendering and optional libraries",
    )
    def test_real_page_previews_are_not_mislabeled_as_visual_review(self):
        from reportlab.pdfgen.canvas import Canvas

        path = self.root / "report.pdf"
        canvas = Canvas(str(path))
        for number in range(2):
            canvas.drawString(72, 720, "Evidence page " + str(number + 1))
            canvas.showPage()
        canvas.save()
        out = self.root / "preview"
        code, value = self.run_cli("knowledge", "inspect-layout", path, "--out", out)
        self.assertEqual(code, 0, value)
        self.assertEqual(value["pages"], 2)
        self.assertEqual(len(list(out.glob("page-*.png"))), 2)
        self.assertEqual(value["visual_review"], "not_performed")
        code, value = self.run_cli("knowledge", "inspect-layout", path, "--out", out)
        self.assertEqual(code, 1, value)


if __name__ == "__main__":
    unittest.main()
