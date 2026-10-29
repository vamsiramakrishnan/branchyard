"""Real library regressions for evidence-preserving table extraction and analysis. No models."""

import csv
import hashlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "sdk/python"))

from branchyard_artifacts import analysis, tables  # noqa: E402
from branchyard_artifacts.common import ArtifactError, Unavailable  # noqa: E402

PANDAS = importlib.util.find_spec("pandas") is not None
PINT = importlib.util.find_spec("pint") is not None
PDF = importlib.util.find_spec("pdfplumber") is not None and importlib.util.find_spec("reportlab") is not None
OFFICE = all(importlib.util.find_spec(name) is not None for name in ["docx", "openpyxl"])


class Workspace(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def spec(self, value, name="spec.json"):
        path = self.root / name
        path.write_text(json.dumps(value), encoding="utf-8")
        return path


@unittest.skipUnless(PANDAS, "install the explicit analysis runtime to test real pandas")
class CleaningTests(Workspace):
    def source(self, text):
        source = self.root / "sales.csv"
        source.write_text(text, encoding="utf-8")
        return source

    def test_preserves_identifiers_literal_na_and_source_with_default_policy(self):
        source = self.source("account_id,region,amount\n00012,NA,1.20\n00013, East ,\n")
        original = source.read_bytes()
        out = self.root / "cleaned"
        result = analysis.clean(self.spec({"schema_version": 1, "input": {"path": source.name}}), out)
        self.assertEqual(result["status"], "passed")
        self.assertEqual(source.read_bytes(), original)
        with (out / "data.csv").open(newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows, [["account_id", "region", "amount"], ["00012", "NA", "1.20"], ["00013", " East ", ""]])
        receipt = json.loads((out / "receipt.json").read_text())
        self.assertEqual(receipt["source"]["sha256"], hashlib.sha256(original).hexdigest())
        self.assertEqual(receipt["output"]["sha256"], hashlib.sha256((out / "data.csv").read_bytes()).hexdigest())
        self.assertEqual(receipt["null_counts"], {"account_id": 0, "region": 0, "amount": 0})

    def test_explicit_trim_null_numeric_and_duplicate_policy_records_actual_counts(self):
        source = self.source("account_id,region,amount\n00012, East , 12 \n00012,West,99\n00013,NA,missing\n")
        value = {
            "schema_version": 1,
            "input": {"path": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest()},
            "trim": ["region", "amount"],
            "null_tokens": ["missing"],
            "types": {"amount": "integer", "region": "string"},
            "deduplicate": {"columns": ["account_id"], "keep": "first"},
        }
        result = analysis.clean(self.spec(value), self.root / "cleaned")
        self.assertEqual(
            result["counts"],
            {"input_rows": 3, "output_rows": 2, "trimmed_cells": 2, "nullified_cells": 1, "duplicates_removed": 1},
        )
        self.assertEqual(
            (self.root / "cleaned/data.csv").read_text(), "account_id,region,amount\n00012,East,12\n00013,NA,\n"
        )
        self.assertEqual(result["null_counts"]["amount"], 1)

    def test_numeric_conversion_rejects_invalid_nonfinite_and_fractional_integers(self):
        for text, kind in [
            ("invalid", "number"),
            ("inf", "number"),
            ("3.5", "integer"),
            ("9223372036854775808", "integer"),
        ]:
            with self.subTest(text=text):
                source = self.source("id,amount\n00012," + text + "\n")
                spec = self.spec({"schema_version": 1, "input": {"path": source.name}, "types": {"amount": kind}})
                with self.assertRaises((ArtifactError, ValueError, TypeError)):
                    analysis.clean(spec, self.root / "cleaned")
                self.assertFalse((self.root / "cleaned").exists())

    def test_integer_conversion_preserves_values_beyond_binary_float_precision(self):
        source = self.source("id,amount\n00012,9007199254740993\n00013,3.0\n")
        spec = self.spec({"schema_version": 1, "input": {"path": source.name}, "types": {"amount": "integer"}})
        analysis.clean(spec, self.root / "cleaned")
        self.assertEqual((self.root / "cleaned/data.csv").read_text(), "id,amount\n00012,9007199254740993\n00013,3\n")

    def test_duplicate_keep_last_and_error_are_explicit(self):
        source = self.source("id,amount\n00012,1\n00012,2\n")
        value = {
            "schema_version": 1,
            "input": {"path": source.name},
            "deduplicate": {"columns": ["id"], "keep": "last"},
        }
        result = analysis.clean(self.spec(value), self.root / "last")
        self.assertEqual(result["counts"]["duplicates_removed"], 1)
        self.assertEqual((self.root / "last/data.csv").read_text(), "id,amount\n00012,2\n")
        value["deduplicate"]["keep"] = "error"
        with self.assertRaisesRegex(ArtifactError, "Duplicate rows"):
            analysis.clean(self.spec(value), self.root / "error")
        self.assertFalse((self.root / "error").exists())

    def test_input_headers_and_ragged_records_are_rejected_without_pandas_index_inference(self):
        for text in ["id,id\n1,2\n", "id,\n1,2\n", "id,amount\n1,2,3\n", "id,amount\n1\n", "id,amount\n\n"]:
            with self.subTest(text=text):
                source = self.source(text)
                with self.assertRaises(ArtifactError):
                    analysis.clean(self.spec({"schema_version": 1, "input": {"path": source.name}}), self.root / "bad")
                self.assertFalse((self.root / "bad").exists())

    def test_closed_specs_reject_wrong_hash_unknown_columns_links_and_limits(self):
        source = self.source("id,amount\n00012,1\n")
        (self.root / "link.csv").symlink_to(source)
        base = {"schema_version": 1, "input": {"path": source.name}}
        variants = [
            {**base, "unknown": True},
            {**base, "input": {"path": source.name, "sha256": "0" * 64}},
            {**base, "input": {"path": "../sales.csv"}},
            {**base, "input": {"path": "link.csv"}},
            {**base, "trim": ["missing"]},
            {**base, "trim": ["id", "id"]},
            {**base, "types": {"id": "guess"}},
            {**base, "null_tokens": [None]},
            {**base, "deduplicate": {"columns": [], "keep": "first"}},
        ]
        for value in variants:
            with self.subTest(value=value), self.assertRaises(ArtifactError):
                analysis.clean(self.spec(value), self.root / "bad")
        with mock.patch.object(analysis, "MAX_INPUT", 4), self.assertRaisesRegex(ArtifactError, "64 MiB"):
            analysis.clean(self.spec(base), self.root / "bad")
        with mock.patch.object(analysis, "MAX_ROWS", 0), self.assertRaisesRegex(ArtifactError, "100000"):
            analysis.clean(self.spec(base), self.root / "bad")

    def test_existing_destination_and_failed_receipt_leave_no_partial_bundle(self):
        source = self.source("id,amount\n00012,1\n")
        spec = self.spec({"schema_version": 1, "input": {"path": source.name}})
        out = self.root / "existing"
        out.mkdir()
        (out / "keep").write_text("original")
        with self.assertRaisesRegex(ArtifactError, "already exists"):
            analysis.clean(spec, out)
        self.assertEqual((out / "keep").read_text(), "original")
        real_write = Path.write_text

        def fail_receipt(path, *args, **kwargs):
            if path.name == "receipt.json":
                raise OSError("injected full disk")
            return real_write(path, *args, **kwargs)

        with mock.patch.object(Path, "write_text", fail_receipt), self.assertRaisesRegex(OSError, "full disk"):
            analysis.clean(spec, self.root / "failed")
        self.assertFalse((self.root / "failed").exists())
        self.assertFalse(any(path.name.startswith(".failed-") for path in self.root.iterdir()))


@unittest.skipUnless(PINT, "install the explicit analysis runtime to test real Pint")
class CalculationTests(Workspace):
    def value(self):
        return {
            "schema_version": 1,
            "variables": {"trip_distance": {"value": 1500, "unit": "meter"}},
            "expression": "trip_distance",
            "unit": "kilometer",
        }

    def test_native_units_and_arithmetic_are_reopened_and_source_preserved(self):
        value = self.value()
        value["expression"] = "(trip_distance + trip_distance) / 2"
        spec = self.spec(value)
        original = spec.read_bytes()
        out = self.root / "result.json"
        report = analysis.calculate(spec, out)
        self.assertEqual(report["result"], {"value": 1.5, "unit": "kilometer"})
        self.assertEqual(json.loads(out.read_text())["result"], report["result"])
        self.assertEqual(spec.read_bytes(), original)
        self.assertEqual(report["sha256"], hashlib.sha256(out.read_bytes()).hexdigest())
        self.assertEqual(report["semantic_review"], "not_performed")

    def test_illustrative_currency_requires_explicit_basis_and_no_exchange_rate(self):
        value = json.loads((ROOT / "sdk/python/branchyard_artifacts/examples/calculate.json").read_text())
        report = analysis.calculate(self.spec(value), self.root / "money.json")
        self.assertEqual(report["result"], {"value": 31160.0, "unit": "USD"})
        self.assertEqual(report["exchange_rates"], "not_used")
        del value["currency_basis"]
        with self.assertRaisesRegex(Exception, "USD"):
            analysis.calculate(self.spec(value), self.root / "missing.json")
        value["currency_basis"] = {"code": "USD", "basis": "Same nominal USD"}
        value["unit"] = "EUR"
        with self.assertRaisesRegex(Exception, "EUR"):
            analysis.calculate(self.spec(value), self.root / "fx.json")

    def test_dimension_mismatch_division_by_zero_and_nonfinite_result_preserve_existing_output(self):
        import pint

        out = self.root / "result.json"
        out.write_text("previous accepted result")
        for change, exception in [
            ({"unit": "second"}, pint.errors.DimensionalityError),
            ({"expression": "trip_distance / 0"}, ZeroDivisionError),
            ({"expression": "trip_distance * 1e308"}, ArtifactError),
        ]:
            with self.subTest(change=change), self.assertRaises(exception):
                analysis.calculate(self.spec({**self.value(), **change}), out)
            self.assertEqual(out.read_text(), "previous accepted result")

    def test_boundary_rejects_undeclared_variables_code_power_and_invalid_input(self):
        variants = [
            {"expression": "missing + trip_distance"},
            {"expression": "__import__('os')"},
            {"expression": "10 ** 1000000000"},
            {"expression": "sqrt(trip_distance)"},
            {"expression": "'text'"},
            {"variables": {"m": {"value": 1, "unit": "meter"}}, "expression": "m"},
            {"variables": {"trip_distance": {"value": True, "unit": "meter"}}},
            {"assumptions": [None]},
            {"currency_basis": {"code": "USD", "basis": ""}},
            {"unknown": "ignored?"},
        ]
        for variant in variants:
            with self.subTest(variant=variant), self.assertRaises(ArtifactError):
                analysis.calculate(self.spec({**self.value(), **variant}), self.root / "bad.json")
        with mock.patch.object(analysis, "MAX_SPEC", 1), self.assertRaisesRegex(ArtifactError, "256 KiB"):
            analysis.calculate(self.spec(self.value()), self.root / "bad.json")
        with self.assertRaises(ArtifactError):
            spec = self.spec(self.value())
            analysis.calculate(spec, spec)


class TableBoundaryTests(Workspace):
    def test_pdf_extraction_refuses_an_unsupported_python_without_loading_an_old_parser(self):
        path = self.root / "source.pdf"
        path.write_bytes(b"not parsed on Python 3.9")
        with (
            mock.patch.object(tables.sys, "version_info", (3, 9, 21)),
            self.assertRaisesRegex(Unavailable, "Python 3.10"),
        ):
            tables.extract(path, self.root / "out.json")
        self.assertFalse((self.root / "out.json").exists())

    def test_preserve_source_and_refuse_links_unsupported_types_or_oversized_input(self):
        path = self.root / "source.txt"
        path.write_text("original")
        with self.assertRaisesRegex(ArtifactError, "preserve"):
            tables.extract(path, path)
        with self.assertRaisesRegex(ArtifactError, "supports"):
            tables.extract(path, self.root / "out.json")
        link = self.root / "link.txt"
        link.symlink_to(path)
        with self.assertRaisesRegex(ArtifactError, "symlink"):
            tables.extract(link, self.root / "out.json")
        with mock.patch.object(tables, "MAX_BYTES", 1), self.assertRaisesRegex(ArtifactError, "64 MiB"):
            tables.extract(path, self.root / "out.json")


@unittest.skipUnless(OFFICE, "install the explicit artifact runtime to inspect real Office files")
class OfficeTableTests(Workspace):
    def test_docx_table_cells_and_locator_match_the_actual_document(self):
        from docx import Document

        source = self.root / "source.docx"
        document = Document()
        table = document.add_table(rows=2, cols=2)
        for row, values in enumerate([["Account", "Amount"], ["00012", "12"]]):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        document.save(source)
        report = tables.extract(source, self.root / "tables.json")
        self.assertEqual(report["table_count"], 1)
        self.assertEqual(
            report["tables"][0], {"locator": {"table": 1}, "rows": [["Account", "Amount"], ["00012", "12"]]}
        )
        self.assertEqual(report["extraction_review"], "not_performed")

    def test_xlsx_extracts_native_tables_preserving_formulas_without_cached_value_claims(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table

        source = self.root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Account", "Amount"])
        sheet.append(["00012", "=1+2"])
        sheet.add_table(Table(displayName="SalesData", ref="A1:B2"))
        workbook.save(source)
        report = tables.extract(source, self.root / "tables.json")
        self.assertEqual(report["table_count"], 1)
        self.assertEqual(report["tables"][0]["rows"], [["Account", "Amount"], ["00012", "=1+2"]])
        self.assertEqual(report["tables"][0]["locator"], {"sheet": "Sales", "table": "SalesData", "range": "A1:B2"})
        self.assertIn("not evaluated", report["warnings"][0])


@unittest.skipUnless(PDF, "PDF table extraction requires its optional modern-Python runtime")
class PdfTableTests(Workspace):
    def pdf(self, table=False, image=False):
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas

        source = self.root / "source.pdf"
        document = canvas.Canvas(str(source))
        if table:
            for y in [700, 720, 740]:
                document.line(50, y, 250, y)
            for x in [50, 150, 250]:
                document.line(x, 700, x, 740)
            for x, y, text in [(55, 725, "Account"), (155, 725, "Amount"), (55, 705, "00012"), (155, 705, "12")]:
                document.drawString(x, y, text)
            document.showPage()
        if image:
            from PIL import Image

            buffer = io.BytesIO()
            Image.new("RGB", (20, 20), "white").save(buffer, format="PNG")
            buffer.seek(0)
            document.drawImage(ImageReader(buffer), 50, 600, 100, 100)
        else:
            document.drawString(50, 700, "A prose-only page without a table")
        document.showPage()
        document.save()
        return source

    def test_pdf_table_bbox_and_page_location_use_pdfplumber_and_bound_source_bytes(self):
        source = self.pdf(table=True)
        report = tables.extract(source, self.root / "tables.json")
        self.assertEqual(report["table_count"], 1)
        self.assertEqual(report["pages_attempted"], 2)
        self.assertEqual(report["tables"][0]["locator"], {"page": 1, "table": 1})
        self.assertEqual(report["tables"][0]["rows"], [["Account", "Amount"], ["00012", "12"]])
        self.assertEqual(len(report["tables"][0]["bbox"]), 4)
        self.assertEqual(report["source"]["sha256"], hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertEqual(report["extraction_status"], "detected")

    def test_prose_no_table_and_image_ocr_requirement_are_distinct(self):
        report = tables.extract(self.pdf(), self.root / "prose.json")
        self.assertEqual((report["status"], report["extraction_status"]), ("passed", "no_tables_detected"))
        report = tables.extract(self.pdf(image=True), self.root / "scan.json")
        self.assertEqual((report["status"], report["extraction_status"]), ("unavailable", "ocr_required"))
        self.assertEqual(report["suspected_scan_pages"], [1])
        report = tables.extract(self.pdf(table=True, image=True), self.root / "mixed.json")
        self.assertEqual((report["status"], report["extraction_status"]), ("passed", "partial_ocr_required"))
        self.assertEqual(report["suspected_scan_pages"], [2])

    def test_table_count_and_report_byte_bounds_leave_existing_output_unchanged(self):
        source = self.pdf(table=True)
        out = self.root / "out.json"
        out.write_text("accepted")
        for constant in ["MAX_PAGES", "MAX_TABLES", "MAX_CELLS", "MAX_OUTPUT"]:
            with (
                self.subTest(constant=constant),
                mock.patch.object(tables, constant, 0),
                self.assertRaises(ArtifactError),
            ):
                tables.extract(source, out)
            self.assertEqual(out.read_text(), "accepted")


if __name__ == "__main__":
    unittest.main()
